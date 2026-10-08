"""网页附件存储与读取；引擎只保存引用和文本，图片在模型请求前装配。"""
from base64 import b64encode
from datetime import datetime, timezone
from io import BytesIO
import json
from pathlib import Path
import re
import shutil
import threading
from uuid import uuid4
from xml.etree import ElementTree
from zipfile import ZipFile

MAX_BYTES = 20 * 1024 * 1024
MAX_FILES = 6
MAX_TEXT = 60_000
MAX_CONTEXT = 120_000
MAX_AUDIO_SECONDS = 120
IMAGE_TYPES = {".jpg", ".jpeg", ".png", ".webp", ".gif"}
AUDIO_TYPES = {".mp3", ".wav", ".m4a", ".webm", ".ogg", ".flac", ".mp4", ".aac"}
TEXT_TYPES = {".txt", ".md", ".csv", ".json"}
FILE_TYPES = TEXT_TYPES | {".pdf", ".docx", ".xlsx"}
ATTACHMENT_ID = re.compile(r"^[a-f0-9]{32}$")
REFERENCE = re.compile(r'^<attachment id="([a-f0-9]{32})" filename="')


class AttachmentError(ValueError):
    pass


class SpeechTranscriber:
    """按需载入 CPU 量化模型；音频留在本应用，不发送给语言模型供应商。"""
    def __init__(self, model_dir):
        self.model_dir = model_dir
        self.model = None
        self.lock = threading.Lock()

    @property
    def ready(self):
        return all((self.model_dir / name).is_file() for name in
                   ("model.bin", "config.json", "tokenizer.json"))

    def transcribe(self, path):
        if not self.ready:
            raise AttachmentError("语音识别模型尚未安装，请先使用文字输入。")
        from faster_whisper import WhisperModel
        from faster_whisper.audio import decode_audio
        import av
        with self.lock:
            with av.open(str(path)) as container:
                if container.duration and container.duration / av.time_base > MAX_AUDIO_SECONDS:
                    raise AttachmentError("录音最长支持 2 分钟，请截短后重试。")
            samples = decode_audio(str(path), sampling_rate=16000)
            duration = len(samples) / 16000
            if duration > MAX_AUDIO_SECONDS:
                raise AttachmentError("录音最长支持 2 分钟，请截短后重试。")
            if duration < 0.3:
                raise AttachmentError("录音太短，请重新录制。")
            if self.model is None:
                self.model = WhisperModel(str(self.model_dir), device="cpu", compute_type="int8",
                                          cpu_threads=2, local_files_only=True)
            # 工作台默认中文输入；短录音自动探测曾将中文误判为泰语。
            segments, info = self.model.transcribe(samples, language="zh", task="transcribe",
                                                   beam_size=3, vad_filter=True,
                                                   condition_on_previous_text=False)
            transcript = "".join(segment.text for segment in segments).strip()
            if not transcript or info.language_probability < 0.3:
                raise AttachmentError("没有识别到清晰的讲话，请重新录制或直接输入文字。")
            return transcript, round(duration, 1)


class AttachmentStore:
    def __init__(self, runtime_dir, asr_model_dir=None):
        self.root = runtime_dir / "attachments"
        self.root.mkdir(parents=True, exist_ok=True)
        self.speech = SpeechTranscriber(asr_model_dir or runtime_dir / "models/faster-whisper-base")

    def get(self, identifier):
        if not ATTACHMENT_ID.fullmatch(identifier):
            raise AttachmentError("附件不存在。")
        try:
            return json.loads((self.root / identifier / "metadata.json").read_text(encoding="utf-8"))
        except (OSError, ValueError):
            raise AttachmentError("附件不存在，请重新上传。") from None

    def public(self, item):
        result = {key: item[key] for key in ("id", "name", "kind", "size", "mime", "warning",
                                           "text_length", "duration", "transcript", "created_at") if key in item}
        if item["kind"] == "file":
            result["preview_text"] = item["markdown"][:3000]
        return result

    def path(self, identifier, preview=False):
        item = self.get(identifier)
        return self.root / identifier / ("preview.jpg" if preview and item["kind"] == "image" else "original")

    def save(self, name, content):
        name = Path((name or "附件").replace("\\", "/")).name[:180]
        suffix = Path(name).suffix.lower()
        if suffix not in IMAGE_TYPES | AUDIO_TYPES | FILE_TYPES:
            raise AttachmentError("支持 JPG、PNG、WebP、GIF、常见录音、PDF、Word、Excel 和文本文件。")
        if not content:
            raise AttachmentError("这个文件是空的，请检查后重试。")
        if len(content) > MAX_BYTES:
            raise AttachmentError("单个附件不能超过 20 MB。")
        identifier = uuid4().hex
        folder = self.root / identifier
        folder.mkdir(mode=0o700)
        item = {"id": identifier, "name": name, "size": len(content), "warning": "",
                "created_at": datetime.now(timezone.utc).isoformat()}
        try:
            (folder / "original").write_bytes(content)
            if suffix in IMAGE_TYPES:
                self._image(folder, content, item)
            elif suffix in AUDIO_TYPES:
                item.update(kind="audio", mime={".mp3": "audio/mpeg", ".m4a": "audio/mp4",
                    ".mp4": "audio/mp4", ".wav": "audio/wav", ".webm": "audio/webm",
                    ".ogg": "audio/ogg", ".flac": "audio/flac", ".aac": "audio/aac"}[suffix])
                transcript, duration = self.speech.transcribe(folder / "original")
                item.update(transcript=transcript, duration=duration, text_length=len(transcript),
                            markdown=transcript, formatter="speech-to-text")
            else:
                value, warning = self._document(suffix, content)
                if not value.strip():
                    raise AttachmentError("没有读到文字。扫描版 PDF 请上传页面截图，表格请检查是否有内容。")
                item.update(kind="file", mime="application/octet-stream", formatter=suffix[1:],
                            text_length=min(len(value), MAX_TEXT), markdown=value[:MAX_TEXT], warning=warning)
                if len(value) > MAX_TEXT:
                    item["warning"] = "内容较长，本次只读取前 60,000 字；需要后续部分时请拆分文件。"
            (folder / "metadata.json").write_text(json.dumps(item, ensure_ascii=False), encoding="utf-8")
            for file in folder.iterdir():
                file.chmod(0o600)
            return self.public(item)
        except AttachmentError:
            shutil.rmtree(folder)
            raise
        except Exception:
            shutil.rmtree(folder)
            raise AttachmentError("无法读取这个附件，请检查文件是否损坏、加密，或换一种格式。") from None

    def _image(self, folder, content, item):
        from PIL import Image, ImageOps
        with Image.open(BytesIO(content)) as original:
            if original.format not in {"JPEG", "PNG", "WEBP", "GIF"} or original.width * original.height > 30_000_000:
                raise AttachmentError("图片格式不支持或尺寸过大，请转成 JPG、PNG 后重试。")
            image = ImageOps.exif_transpose(original).convert("RGB")
            image.thumbnail((3072, 3072))
            image.save(folder / "preview.jpg", format="JPEG", quality=93)
            item.update(kind="image", mime="image/jpeg", formatter="image", text_length=0,
                        markdown="图片附件；请读取随附的视觉内容。")
            if getattr(original, "n_frames", 1) > 1:
                item["warning"] = "动态图本次读取第一帧。"

    def _document(self, suffix, content):
        warning = ""
        if suffix in TEXT_TYPES:
            try:
                value = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                value = content.decode("gb18030")
            if "\x00" in value:
                raise AttachmentError("这个文件不像文本文件，请检查格式。")
            return value, warning
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise AttachmentError("暂不支持加密 PDF，请解密后上传。")
            pieces, length = [], 0
            for number, page in enumerate(reader.pages[:60], 1):
                piece = f"\n--- 第 {number} 页 ---\n" + (page.extract_text() or "")
                pieces.append(piece)
                length += len(piece)
                if length > MAX_TEXT:
                    break
            if len(reader.pages) > 60:
                warning = "本次最多读取前 60 页；需要后续页面时请拆分 PDF。"
            value = "".join(pieces)
            if not re.sub(r"--- 第 \d+ 页 ---", "", value).strip():
                value = ""
            return value, warning
        if suffix == ".docx":
            with ZipFile(BytesIO(content)) as archive:
                entry = archive.getinfo("word/document.xml")
                if entry.file_size > 30 * 1024 * 1024:
                    raise AttachmentError("Word 文档内容过大，请拆分后上传。")
                tree = ElementTree.fromstring(archive.read(entry))
                namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
                lines = ["".join(node.text or "" for node in paragraph.iter(namespace + "t"))
                         for paragraph in tree.iter(namespace + "p")]
                return "\n".join(lines), warning
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            pieces, length, rows = [], 0, 0
            for sheet in workbook:
                pieces.append("\n## 工作表：" + sheet.title + "\n")
                for row in sheet.iter_rows(values_only=True):
                    if not any(value is not None for value in row):
                        continue
                    line = " | ".join("" if value is None else str(value) for value in row) + "\n"
                    pieces.append(line)
                    length += len(line)
                    rows += 1
                    if length > MAX_TEXT or rows >= 3000:
                        return "".join(pieces), "表格较大，本次最多读取 3,000 个非空行或前 60,000 字，请按需拆分。"
            return "".join(pieces) if rows else "", warning
        finally:
            workbook.close()

    def resolve(self, identifiers):
        items = [self.get(identifier) for identifier in dict.fromkeys(identifiers)]
        if sum(len(item["markdown"]) for item in items) > MAX_CONTEXT:
            raise AttachmentError("附件文字总量超过 120,000 字，请减少附件或拆分发送。")
        if sum(self.path(item["id"], preview=True).stat().st_size for item in items
               if item["kind"] == "image") > 24 * 1024 * 1024:
            raise AttachmentError("图片总量较大，请减少图片后发送。")
        return [{"attachment_id": item["id"], "filename": item["name"], "formatter": item["formatter"],
                 "markdown": item["markdown"] + ("\n读取范围：" + item["warning"] if item["warning"] else "")}
                for item in items]

    def from_content(self, blocks):
        result = []
        for block in blocks if isinstance(blocks, list) else []:
            match = REFERENCE.match(str(block.get("text", ""))) if isinstance(block, dict) else None
            if match:
                try:
                    result.append(self.public(self.get(match[1])))
                except AttachmentError:
                    result.append({"id": match[1], "name": "附件已不可用", "kind": "file", "missing": True})
        return result

    def image_block(self, identifier, protocol):
        payload = b64encode(self.path(identifier, preview=True).read_bytes()).decode("ascii")
        if protocol == "claude":
            return {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": payload}}
        return {"type": "image_url", "image_url": {"url": "data:image/jpeg;base64," + payload}}
